"""
GENOPOISK CRM — восстановление дат после первичного импорта Отчеты_2025-2026.xlsx.

Первичный импорт (app/historical_import.py) намеренно оставлял пустыми все
даты, которые Excel хранил не как настоящую дату (float-огрызки вроде 10.04,
либо свободный текст вроде "перезабор 3.03") — раздел "не угадывать".

Этот скрипт — отдельный, обдуманный проход поверх уже импортированных
данных, по запросу пользователя после ревью. Правило восстановления:

  Для числового огрызка ДД.ММ ищем ближайшую настоящую дату той же колонки
  ДО и ПОСЛЕ данной строки (по номеру строки Excel). Перебираем года, которые
  реально встречаются в этой колонке. Если РОВНО один год даёт дату, которая
  укладывается между "до" и "после" по календарю — восстанавливаем именно её
  и год восстановления логируем в audit_log. Если ни один год не подходит,
  или подходят несколько (неоднозначность), или соседних дат вообще нет —
  НЕ восстанавливаем, оставляем пустым.

  Для свободного текста (не число) — восстановление невозможно в принципе;
  вместо этого исходное значение дописывается в comments теста, чтобы оно
  не потерялось молча.
"""
from __future__ import annotations

import datetime
import sqlite3
from pathlib import Path

import openpyxl

from app import db as dbmod
from app import repository as repo


def get_col_data(wb, sheet, col_idx):
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    data = [r for r in rows[1:] if not all(c is None for c in r)]
    out = {}
    for i, r in enumerate(data, start=2):
        out[i] = r[col_idx]
    return out


def parse_float_ddmm(v):
    s = f"{v:.2f}"
    day_s, frac = s.split(".")
    day, month = int(day_s), int(frac)
    if not (1 <= day <= 31 and 1 <= month <= 12):
        return None
    return day, month


# (sheet, col_idx, gns_col_idx, target_field) — target_field на тесте, куда пишем результат
COLUMN_TARGETS = [
    ("WGS", 7, 0, "sample_received_at"),
    ("WGS", 8, 0, "client_issued_at"),
    ("SEM", 5, 0, "sample_received_at"),
    ("SEM", 6, 0, "client_issued_at"),
    ("BASH", 5, 0, "sample_received_at"),
    ("ROM", 6, 0, "sample_received_at"),
]

# номера SEM, которым при первичном импорте присвоен суффикс -мтДНК (для второго
# вхождения M-1..M-5) или переномерованы (SE2->SEZ) — нужно для сопоставления
# исходной строки Excel с итоговым gns_number в БД.
SEM_RENUMBER_BY_ROW = {36: "M-1-мтДНК", 37: "M-2-мтДНК", 38: "M-3-мтДНК",
                        39: "M-4-мтДНК", 40: "M-5-мтДНК", 53: "SEZ"}


def resolve_gns_for_row(sheet: str, excel_row: int, raw_number: str) -> str:
    if sheet == "SEM" and excel_row in SEM_RENUMBER_BY_ROW:
        return SEM_RENUMBER_BY_ROW[excel_row]
    if sheet == "Общий" and raw_number.strip().upper().startswith("GNPKS"):
        return raw_number.strip().replace("GNPKS", "GNPSK", 1)
    return raw_number.strip()


def run_remediation(conn: sqlite3.Connection, xlsx_path: str, *, user_name: str = "date-remediation-2026-08-13"):
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)

    restored_count = 0
    text_preserved_count = 0
    unresolved_count = 0
    restored_log = []
    unresolved_log = []
    text_log = []

    for sheet, col_idx, gns_col_idx, target_field in COLUMN_TARGETS:
        col = get_col_data(wb, sheet, col_idx)
        gns_col = get_col_data(wb, sheet, gns_col_idx)

        anchors = sorted(
            (i, v.date() if isinstance(v, datetime.datetime) else v)
            for i, v in col.items() if isinstance(v, (datetime.date, datetime.datetime))
        )
        years_seen = sorted({d.year for _, d in anchors})

        for i, v in col.items():
            raw_gns = gns_col.get(i)
            if raw_gns is None:
                continue
            gns = resolve_gns_for_row(sheet, i, str(raw_gns))

            if isinstance(v, (datetime.date, datetime.datetime)) or v is None:
                continue

            row = conn.execute("SELECT test_id, comments, %s FROM tests WHERE gns_number = ?" % target_field, (gns,)).fetchone()
            if row is None:
                continue  # строка была удалена/пропущена при первичном импорте

            if isinstance(v, float):
                dm = parse_float_ddmm(v)
                if dm is None:
                    unresolved_count += 1
                    unresolved_log.append((sheet, i, gns, v, "не похоже на ДД.ММ"))
                    continue
                day, month = dm
                before = max((a for a in anchors if a[0] < i), key=lambda a: a[0], default=None)
                after = min((a for a in anchors if a[0] > i), key=lambda a: a[0], default=None)

                candidates = []
                for y in years_seen:
                    try:
                        cand = datetime.date(y, month, day)
                    except ValueError:
                        continue
                    ok_before = (before is None) or (cand >= before[1])
                    ok_after = (after is None) or (cand <= after[1])
                    if ok_before and ok_after:
                        candidates.append(cand)

                if len(candidates) == 1:
                    new_date = candidates[0]
                    if target_field == "client_issued_at":
                        new_val = new_date.strftime("%Y-%m-%dT00:00:00")
                    else:
                        new_val = new_date.isoformat()
                    conn.execute("BEGIN IMMEDIATE")
                    conn.execute(f"UPDATE tests SET {target_field} = ?, updated_at = ? WHERE test_id = ?",
                                 (new_val, repo.now_iso(), row["test_id"]))
                    repo.write_audit(
                        conn, user_name=user_name, entity="test", entity_id=row["test_id"],
                        field=target_field, old_value=None, new_value=new_val,
                        reason=f"восстановлена дата из '{v}' по контексту листа {sheet} "
                               f"(единственный год {new_date.year}, укладывается между соседними датами)",
                    )
                    conn.execute("COMMIT")
                    restored_count += 1
                    restored_log.append((sheet, i, gns, v, new_val))
                else:
                    unresolved_count += 1
                    reason = "несколько лет подходят" if len(candidates) > 1 else "ни один год не укладывается между соседними датами"
                    unresolved_log.append((sheet, i, gns, v, reason))
            else:
                # свободный текст — сохранить в comments, если ещё не сохранён
                text_repr = str(v)
                current_comment = row["comments"] or ""
                marker = f"[исходное значение поля даты: {text_repr!r}]"
                if marker not in current_comment:
                    new_comment = (current_comment + "; " if current_comment else "") + marker
                    conn.execute("BEGIN IMMEDIATE")
                    conn.execute("UPDATE tests SET comments = ?, updated_at = ? WHERE test_id = ?",
                                 (new_comment, repo.now_iso(), row["test_id"]))
                    repo.write_audit(
                        conn, user_name=user_name, entity="test", entity_id=row["test_id"],
                        field="comments", old_value=current_comment, new_value=new_comment,
                        reason="сохранение нераспознаваемого текстового значения даты (не восстанавливается)",
                    )
                    conn.execute("COMMIT")
                text_preserved_count += 1
                text_log.append((sheet, i, gns, text_repr))

    wb.close()
    return {
        "restored_count": restored_count,
        "text_preserved_count": text_preserved_count,
        "unresolved_count": unresolved_count,
        "restored_log": restored_log,
        "unresolved_log": unresolved_log,
        "text_log": text_log,
    }


def analyze_remediation(conn: sqlite3.Connection, xlsx_path: str):
    """Чистый анализ БЕЗ записи в БД — для отчётов. Логика идентична run_remediation,
    но не выполняет UPDATE/audit_log, только вычисляет, что было бы сделано."""
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    restored_log, unresolved_log, text_log = [], [], []

    for sheet, col_idx, gns_col_idx, target_field in COLUMN_TARGETS:
        col = get_col_data(wb, sheet, col_idx)
        gns_col = get_col_data(wb, sheet, gns_col_idx)
        anchors = sorted(
            (i, v.date() if isinstance(v, datetime.datetime) else v)
            for i, v in col.items() if isinstance(v, (datetime.date, datetime.datetime))
        )
        years_seen = sorted({d.year for _, d in anchors})

        for i, v in col.items():
            raw_gns = gns_col.get(i)
            if raw_gns is None or isinstance(v, (datetime.date, datetime.datetime)) or v is None:
                continue
            gns = resolve_gns_for_row(sheet, i, str(raw_gns))
            row = conn.execute("SELECT test_id FROM tests WHERE gns_number = ?", (gns,)).fetchone()
            if row is None:
                continue

            if isinstance(v, float):
                dm = parse_float_ddmm(v)
                if dm is None:
                    unresolved_log.append((sheet, i, gns, v, "не похоже на ДД.ММ"))
                    continue
                day, month = dm
                before = max((a for a in anchors if a[0] < i), key=lambda a: a[0], default=None)
                after = min((a for a in anchors if a[0] > i), key=lambda a: a[0], default=None)
                candidates = []
                for y in years_seen:
                    try:
                        cand = datetime.date(y, month, day)
                    except ValueError:
                        continue
                    if ((before is None) or (cand >= before[1])) and ((after is None) or (cand <= after[1])):
                        candidates.append(cand)
                if len(candidates) == 1:
                    new_date = candidates[0]
                    new_val = new_date.strftime("%Y-%m-%dT00:00:00") if target_field == "client_issued_at" else new_date.isoformat()
                    restored_log.append((sheet, i, gns, v, new_val))
                else:
                    reason = "несколько лет подходят" if len(candidates) > 1 else "ни один год не укладывается между соседними датами"
                    unresolved_log.append((sheet, i, gns, v, reason))
            else:
                text_log.append((sheet, i, gns, str(v)))

    wb.close()
    return {"restored_log": restored_log, "unresolved_log": unresolved_log, "text_log": text_log}


if __name__ == "__main__":
    import sys
    db_path = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("/home/claude/genopoisk_crm_migration/genopoisk.db")
    xlsx_path = sys.argv[2] if len(sys.argv) > 2 else "/mnt/user-data/uploads/Отчеты_2025-2026.xlsx"

    conn = dbmod.connect(db_path)
    result = run_remediation(conn, xlsx_path)

    print(f"Восстановлено автоматически: {result['restored_count']}")
    for x in result["restored_log"]:
        print("  ", x)
    print(f"\nСохранён исходный текст в comments (не восстановлено): {result['text_preserved_count']}")
    for x in result["text_log"]:
        print("  ", x)
    print(f"\nНе восстановлено (неоднозначно/несовместимо): {result['unresolved_count']}")
    for x in result["unresolved_log"]:
        print("  ", x)

    conn.close()
