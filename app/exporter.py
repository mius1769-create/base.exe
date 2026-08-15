"""
GENOPOISK CRM — экспорт заказов/тестов в Excel (задача 4 финального этапа).

Одна строка = один тест (с данными его заказа) — так же, как отображается
в основном списке приложения, чтобы выгрузка была узнаваемой.
"""
from __future__ import annotations

import sqlite3
from typing import Optional

import csv

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from . import repository as repo
from . import product_catalog as pcat

COLUMNS = [
    ("gns_number", "GNS"),
    ("order_no", "№ заказа"),
    ("customer_name", "ФИО"),
    ("contacts", "Контакты"),
    ("test_type", "Тип теста"),          # коммерческое название (напр. "Y Базовый")
    ("_technical_code", "Технический код"),  # исходный код из БД (напр. "Y50") — для сверки/отладки персоналом
    ("report_option", "Комплектация"),
    ("order_status", "Статус заказа"),
    ("sample_received_at", "Дата получения образца"),
    ("lab_sent_at", "Дата передачи в лабораторию"),
    ("lab_profile_received_at", "Дата профиля лаборатории"),
    ("operator_started_at", "Взято в работу"),
    ("operator_name", "Оператор"),
    ("client_issued_at", "Выдан клиенту"),
    ("deadline_date", "Дедлайн"),
    ("result_y", "Результат Y"),
    ("result_mt", "Результат мтДНК"),
    ("comments", "Комментарий"),
    ("is_historical", "Историческая запись"),
]


def export_orders_to_excel(
    conn: sqlite3.Connection,
    file_path: str,
    *,
    search: Optional[str] = None,
    quick_filter: str = "all",
    include_archived: bool = False,
) -> int:
    """
    Выгружает тесты (с данными заказа) в Excel. Использует те же
    search/quick_filter, что и главный экран — можно выгрузить как всё,
    так и текущий отфильтрованный список. Возвращает число выгруженных строк.
    """
    rows = repo.list_tests_with_order(
        conn, search=search, quick_filter=quick_filter, include_archived=include_archived
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Тесты"

    headers = [label for _, label in COLUMNS]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for row in rows:
        values = []
        for field_name, _label in COLUMNS:
            if field_name == "_technical_code":
                v = row["test_type"]
            elif field_name == "test_type":
                v = pcat.get_display_name(row["test_type"])
            elif field_name == "is_historical":
                v = "Да" if row[field_name] else ""
            else:
                v = row[field_name]
            values.append(v if v is not None else "")
        ws.append(values)

    # разумная ширина колонок, чтобы не открывать нечитаемый лист
    widths = {"GNS": 16, "ФИО": 24, "Комментарий": 40, "Контакты": 22}
    for i, (_field, label) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(label, 16)

    ws.freeze_panes = "A2"

    wb.save(file_path)
    return len(rows)


# ---------------------------------------------------------------------
# Экспорт вкладки «Гаплогруппы» (задача: третья вкладка CRM)
#
# Принимает уже подготовленный вызывающим кодом (haplogroups_tab.py)
# список строк как raw-словарей — то же представление, что видно в
# таблице (с учётом текущих фильтров), а не отдельный запрос к БД.
# ---------------------------------------------------------------------

HAPLOGROUP_COLUMNS = [
    ("test_number", "№ теста"),
    ("full_name", "ФИО"),
    ("project_path", "Проект"),
    ("y_dna", "Y-ДНК"),
    ("mt_dna", "mtDNA"),
    ("nevgen_prediction", "Предикция NevGen"),
    ("semargl_prediction", "Предикция Semargl"),
    ("snp_issued", "SNP клиенту"),
    ("comment", "Комментарий"),
]


def export_haplogroups_to_excel(rows: list[dict], file_path: str) -> int:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Гаплогруппы"

    headers = [label for _, label in HAPLOGROUP_COLUMNS]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for row in rows:
        ws.append([row.get(field_name) or "" for field_name, _label in HAPLOGROUP_COLUMNS])

    widths = {"ФИО": 24, "Проект": 24, "Комментарий": 40, "Предикция NevGen": 22, "Предикция Semargl": 22}
    for i, (_field, label) in enumerate(HAPLOGROUP_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(label, 16)

    ws.freeze_panes = "A2"

    wb.save(file_path)
    return len(rows)


def export_haplogroups_to_csv(rows: list[dict], file_path: str) -> int:
    with open(file_path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f, delimiter=";")
        writer.writerow([label for _, label in HAPLOGROUP_COLUMNS])
        for row in rows:
            writer.writerow([row.get(field_name) or "" for field_name, _label in HAPLOGROUP_COLUMNS])
    return len(rows)
