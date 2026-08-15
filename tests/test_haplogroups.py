"""
Тесты вкладки №3 «Гаплогруппы»: схема БД, дерево проектов, цветовая
логика NevGen/Semargl, фильтры, экспорт Excel/CSV.
"""
import os
import sqlite3
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

import pytest

from app import db as dbmod
from app import repository as repo
from app import projects_repo as prepo
from app import haplogroups_repo as hrepo
from app import haplogroup_logic as hlogic
from app import exporter


@pytest.fixture()
def conn(tmp_path):
    c = dbmod.connect(tmp_path / "test.db")
    dbmod.init_db(c)
    yield c
    c.close()


def _make_test(conn, customer_name="Тестовый Клиент", **kwargs):
    order = repo.NewOrderInput(customer_name=customer_name, test_type_codes=["Y50"], **kwargs)
    order_id, test_ids = repo.create_order_with_tests(conn, order)
    gns = conn.execute("SELECT gns_number FROM tests WHERE test_id = ?", (test_ids[0],)).fetchone()[0]
    return gns


# ---------------------------------------------------------------------
# Схема БД — только добавление, ничего не сломано у существующих таблиц
# ---------------------------------------------------------------------

def test_existing_tables_untouched(conn):
    for table in ("orders", "tests", "test_types", "test_events", "audit_log", "settings"):
        conn.execute(f"SELECT * FROM {table} LIMIT 1")  # не падает


def test_projects_and_haplogroups_tables_exist(conn):
    conn.execute("SELECT * FROM projects LIMIT 1")
    conn.execute("SELECT * FROM haplogroups LIMIT 1")


def test_default_projects_seeded(conn):
    names = {r["name"] for r in prepo.list_projects(conn)}
    assert {"Клиенты", "Этнопроекты", "Башкирский проект", "Татарский проект"} <= names
    ethno = conn.execute("SELECT id FROM projects WHERE name = 'Этнопроекты'").fetchone()["id"]
    bashkir = conn.execute("SELECT parent_id FROM projects WHERE name = 'Башкирский проект'").fetchone()
    assert bashkir["parent_id"] == ethno


def test_reserved_fields_present_in_schema(conn):
    cols = {r[1] for r in conn.execute("PRAGMA table_info(haplogroups)").fetchall()}
    for reserved in ("date_prediction", "analyst", "review_status", "date_issued"):
        assert reserved in cols
    for visible in ("test_number", "full_name", "project_id", "y_dna", "mt_dna",
                     "nevgen_prediction", "semargl_prediction", "snp_issued", "comment"):
        assert visible in cols


def test_reserved_fields_not_in_editable_ui_fields():
    assert hrepo.RESERVED_FIELDS.isdisjoint(hrepo.EDITABLE_FIELDS)
    assert "review_status" in hrepo.RESERVED_FIELDS
    assert "analyst" in hrepo.RESERVED_FIELDS
    assert "date_issued" in hrepo.RESERVED_FIELDS
    assert "date_prediction" in hrepo.RESERVED_FIELDS


# ---------------------------------------------------------------------
# Дерево проектов
# ---------------------------------------------------------------------

def test_create_nested_project(conn):
    root_id = prepo.create_project(conn, "Клиенты B2B")
    child_id = prepo.create_project(conn, "VIP-клиенты", root_id)
    tree = prepo.build_tree(conn)
    node = next(n for n in tree if n["id"] == root_id)
    assert node["children"][0]["id"] == child_id


def test_rename_project(conn):
    pid = prepo.create_project(conn, "Старое имя")
    prepo.rename_project(conn, pid, "Новое имя")
    row = conn.execute("SELECT name FROM projects WHERE id = ?", (pid,)).fetchone()
    assert row["name"] == "Новое имя"


def test_move_project_prevents_cycle(conn):
    parent_id = prepo.create_project(conn, "Родитель")
    child_id = prepo.create_project(conn, "Ребёнок", parent_id)
    with pytest.raises(ValueError):
        prepo.move_project(conn, parent_id, child_id)


def test_archive_and_unarchive_project(conn):
    pid = prepo.create_project(conn, "Архивный проект")
    prepo.archive_project(conn, pid)
    assert conn.execute("SELECT is_archived FROM projects WHERE id=?", (pid,)).fetchone()["is_archived"] == 1
    prepo.unarchive_project(conn, pid)
    assert conn.execute("SELECT is_archived FROM projects WHERE id=?", (pid,)).fetchone()["is_archived"] == 0


def test_get_project_path_breadcrumbs(conn):
    ethno = conn.execute("SELECT id FROM projects WHERE name = 'Этнопроекты'").fetchone()["id"]
    bashkir = conn.execute("SELECT id FROM projects WHERE name = 'Башкирский проект'").fetchone()["id"]
    assert prepo.get_project_path(conn, bashkir) == "Этнопроекты › Башкирский проект"
    assert prepo.get_project_path(conn, ethno) == "Этнопроекты"
    assert prepo.get_project_path(conn, None) == ""


def test_get_descendant_ids(conn):
    ethno = conn.execute("SELECT id FROM projects WHERE name = 'Этнопроекты'").fetchone()["id"]
    bashkir = conn.execute("SELECT id FROM projects WHERE name = 'Башкирский проект'").fetchone()["id"]
    tatar = conn.execute("SELECT id FROM projects WHERE name = 'Татарский проект'").fetchone()["id"]
    descendants = prepo.get_descendant_ids(conn, ethno)
    assert {bashkir, tatar} <= descendants


# ---------------------------------------------------------------------
# Цветовая логика NevGen vs Semargl (якорь — подтверждённое поле Y-ДНК,
# а не сравнение root(NevGen) с root(Semargl) друг с другом)
# ---------------------------------------------------------------------

def test_color_green_when_snp_found_in_semargl():
    assert hlogic.calculate_haplo_color(
        "R1a-Z93-Z94", "R1a-Z93-Z94-YP1337", "R1a"
    ) == hlogic.HaploColor.GREEN


def test_color_green_examples_from_review():
    assert hlogic.calculate_haplo_color(
        "I-CTS10228", "I > CTS10228 > Y3120 > PH908", "I2a"
    ) == hlogic.HaploColor.GREEN
    assert hlogic.calculate_haplo_color(
        "R-M458", "R1a > M458 > YP417", "R1a"
    ) == hlogic.HaploColor.GREEN


def test_color_yellow_when_snp_missing_but_y_dna_root_matches_semargl():
    """Ревью-кейс: Y-ДНК=R1a, NevGen=R-Z280, Semargl=R-M458 -> SNP не найден,
    но корень Y-ДНК совпадает с корнем Semargl -> YELLOW."""
    assert hlogic.calculate_haplo_color(
        "R-Z280", "R-M458", "R1a"
    ) == hlogic.HaploColor.YELLOW


def test_color_red_when_snp_missing_and_y_dna_root_differs_from_semargl():
    """Ревью-кейс: Y-ДНК=I2a, Semargl указывает ветвь R1a -> корни не
    совпадают -> RED, независимо от того, что написано в NevGen."""
    assert hlogic.calculate_haplo_color(
        "I-M223", "R1a > M405 > L459", "I2a"
    ) == hlogic.HaploColor.RED


def test_color_nevgen_own_root_is_never_compared_directly_to_semargl():
    """Бывший баг (пример из ревью): NevGen 'R-M458' и Semargl 'R > Z280 >
    CTS1211' имеют один и тот же терсовый корневой токен 'R', и старая
    логика (root(NevGen) == root(Semargl)) давала здесь YELLOW — хотя это
    разные ветви. Новая логика вообще не сравнивает root(NevGen) с
    root(Semargl) друг с другом: корень берётся только из подтверждённого
    Y-ДНК. Если Y-ДНК указывает на другую ветвь ('I2a'), результат — RED,
    несмотря на то, что NevGen и Semargl совпадали бы по старому критерию."""
    assert hlogic.calculate_haplo_color(
        "R-M458", "R > Z280 > CTS1211", "I2a"
    ) == hlogic.HaploColor.RED


def test_color_none_when_any_field_missing():
    assert hlogic.calculate_haplo_color("", "R1a-Z93", "R1a") == hlogic.HaploColor.NONE
    assert hlogic.calculate_haplo_color("R1a-Z93", None, "R1a") == hlogic.HaploColor.NONE
    assert hlogic.calculate_haplo_color("R1a-Z93", "R1a-Z93", "") == hlogic.HaploColor.NONE
    assert hlogic.calculate_haplo_color("R1a-Z93", "R1a-Z93", None) == hlogic.HaploColor.NONE
    assert hlogic.calculate_haplo_color(None, None, None) == hlogic.HaploColor.NONE


def test_color_green_case_insensitive():
    assert hlogic.calculate_haplo_color(
        "r1a-z94", "R1A-Z93-Z94", "r1a"
    ) == hlogic.HaploColor.GREEN


# ---------------------------------------------------------------------
# CRUD записей гаплогрупп + автозаполнение full_name
# ---------------------------------------------------------------------

def test_upsert_creates_record_with_full_name_from_order(conn):
    gns = _make_test(conn, customer_name="Иванов Иван Иванович")
    hap_id = hrepo.upsert_haplogroup(conn, gns, {"y_dna": "R1a", "nevgen_prediction": "R1a-Z93"})
    row = conn.execute("SELECT * FROM haplogroups WHERE id = ?", (hap_id,)).fetchone()
    assert row["full_name"] == "Иванов Иван Иванович"
    assert row["test_number"] == gns
    assert row["y_dna"] == "R1a"


def test_upsert_unknown_field_raises(conn):
    gns = _make_test(conn)
    with pytest.raises(ValueError):
        hrepo.upsert_haplogroup(conn, gns, {"not_a_real_field": "x"})


def test_upsert_unknown_test_number_raises(conn):
    with pytest.raises(ValueError):
        hrepo.upsert_haplogroup(conn, "GNPSK-NOPE", {"y_dna": "R1a"})


def test_upsert_is_idempotent_update_not_duplicate(conn):
    gns = _make_test(conn)
    hrepo.upsert_haplogroup(conn, gns, {"y_dna": "R1a"})
    hrepo.upsert_haplogroup(conn, gns, {"y_dna": "R1b"})
    rows = conn.execute("SELECT * FROM haplogroups WHERE test_number = ?", (gns,)).fetchall()
    assert len(rows) == 1
    assert rows[0]["y_dna"] == "R1b"


def test_upsert_writes_audit_log(conn):
    gns = _make_test(conn)
    hrepo.upsert_haplogroup(conn, gns, {"y_dna": "R1a"}, user_name="analyst1")
    audit = conn.execute("SELECT * FROM audit_log WHERE entity='haplogroup'").fetchall()
    assert len(audit) >= 1


def test_reserved_fields_can_be_written_but_stay_hidden(conn):
    gns = _make_test(conn)
    hrepo.upsert_haplogroup(conn, gns, {"analyst": "Петров", "review_status": "проверено"})
    row = conn.execute("SELECT * FROM haplogroups WHERE test_number=?", (gns,)).fetchone()
    assert row["analyst"] == "Петров"
    assert row["review_status"] == "проверено"


# ---------------------------------------------------------------------
# Фильтры списка
# ---------------------------------------------------------------------

def test_list_haplogroups_filters_by_project(conn):
    ethno = conn.execute("SELECT id FROM projects WHERE name = 'Этнопроекты'").fetchone()["id"]
    bashkir = conn.execute("SELECT id FROM projects WHERE name = 'Башкирский проект'").fetchone()["id"]
    clients = conn.execute("SELECT id FROM projects WHERE name = 'Клиенты'").fetchone()["id"]

    gns1 = _make_test(conn, customer_name="Клиент Один")
    gns2 = _make_test(conn, customer_name="Клиент Два")
    hrepo.upsert_haplogroup(conn, gns1, {"project_id": bashkir})
    hrepo.upsert_haplogroup(conn, gns2, {"project_id": clients})

    # фильтр по родителю "Этнопроекты" должен захватывать вложенный "Башкирский проект"
    rows = hrepo.list_haplogroups(conn, project_id=ethno)
    assert {r["test_number"] for r in rows} == {gns1}

    rows2 = hrepo.list_haplogroups(conn, project_id=clients)
    assert {r["test_number"] for r in rows2} == {gns2}


def test_list_haplogroups_filters_by_text_fields(conn):
    gns1 = _make_test(conn, customer_name="Сидоров Сидор")
    gns2 = _make_test(conn, customer_name="Петров Пётр")
    hrepo.upsert_haplogroup(conn, gns1, {"y_dna": "R1a", "mt_dna": "H1", "snp_issued": "Z93"})
    hrepo.upsert_haplogroup(conn, gns2, {"y_dna": "I2", "mt_dna": "U5", "snp_issued": "L621"})

    assert {r["test_number"] for r in hrepo.list_haplogroups(conn, full_name="Сидоров")} == {gns1}
    assert {r["test_number"] for r in hrepo.list_haplogroups(conn, y_dna="I2")} == {gns2}
    assert {r["test_number"] for r in hrepo.list_haplogroups(conn, mt_dna="H1")} == {gns1}
    assert {r["test_number"] for r in hrepo.list_haplogroups(conn, snp_issued="L621")} == {gns2}
    assert {r["test_number"] for r in hrepo.list_haplogroups(conn, test_number=gns1)} == {gns1}


def test_list_haplogroups_excludes_archived_by_default(conn):
    gns = _make_test(conn)
    hap_id = hrepo.upsert_haplogroup(conn, gns, {"y_dna": "R1a"})
    hrepo.archive_haplogroup(conn, hap_id)
    assert hrepo.list_haplogroups(conn, test_number=gns) == []
    assert len(hrepo.list_haplogroups(conn, test_number=gns, include_archived=True)) == 1


def test_display_full_name_reflects_live_order_data(conn):
    """full_name «из карточки» — список всегда показывает актуальное ФИО заказа,
    даже если снимок в haplogroups.full_name устарел."""
    gns = _make_test(conn, customer_name="Старое ФИО")
    hrepo.upsert_haplogroup(conn, gns, {"y_dna": "R1a"})
    conn.execute("UPDATE orders SET customer_name = 'Новое ФИО' WHERE order_id = "
                 "(SELECT order_id FROM tests WHERE gns_number = ?)", (gns,))
    rows = hrepo.list_haplogroups(conn, test_number=gns)
    assert rows[0]["display_full_name"] == "Новое ФИО"


# ---------------------------------------------------------------------
# Экспорт
# ---------------------------------------------------------------------

def test_export_haplogroups_to_excel(tmp_path, conn):
    gns = _make_test(conn, customer_name="Экспортов Экспорт")
    hrepo.upsert_haplogroup(conn, gns, {"y_dna": "R1a", "nevgen_prediction": "R1a-Z93",
                                         "semargl_prediction": "R1a-Z93-Z94", "comment": "тест"})
    rows = hrepo.list_haplogroups(conn)
    export_rows = [{
        "test_number": r["test_number"], "full_name": r["display_full_name"],
        "project_path": prepo.get_project_path(conn, r["project_id"]),
        "y_dna": r["y_dna"], "mt_dna": r["mt_dna"],
        "nevgen_prediction": r["nevgen_prediction"], "semargl_prediction": r["semargl_prediction"],
        "snp_issued": r["snp_issued"], "comment": r["comment"],
    } for r in rows]

    path = tmp_path / "export.xlsx"
    count = exporter.export_haplogroups_to_excel(export_rows, str(path))
    assert count == 1
    assert path.exists()

    import openpyxl
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    header = [c.value for c in ws[1]]
    assert header == [label for _f, label in exporter.HAPLOGROUP_COLUMNS]
    assert ws.cell(row=2, column=1).value == gns


def test_export_haplogroups_to_csv(tmp_path, conn):
    gns = _make_test(conn)
    hrepo.upsert_haplogroup(conn, gns, {"y_dna": "R1a"})
    export_rows = [{"test_number": gns, "full_name": "X", "project_path": "", "y_dna": "R1a",
                     "mt_dna": "", "nevgen_prediction": "", "semargl_prediction": "",
                     "snp_issued": "", "comment": ""}]
    path = tmp_path / "export.csv"
    count = exporter.export_haplogroups_to_csv(export_rows, str(path))
    assert count == 1
    content = path.read_text(encoding="utf-8-sig")
    assert gns in content
    assert "№ теста" in content


# ---------------------------------------------------------------------
# Схема не сломала существующую бизнес-логику 1/2 вкладок
# ---------------------------------------------------------------------

def test_existing_order_and_event_flow_still_works(conn):
    gns = _make_test(conn)
    test_id = conn.execute("SELECT test_id FROM tests WHERE gns_number=?", (gns,)).fetchone()["test_id"]
    repo.record_event(conn, test_id, "received", operator="оператор")
    events = repo.list_events_for_test(conn, test_id)
    assert any(e["event_type"] == "received" for e in events)


# ---------------------------------------------------------------------
# Доп. проверки по итогам ревью
# ---------------------------------------------------------------------

def test_db_level_unique_constraint_on_test_number(conn):
    """UNIQUE(test_number) в схеме — вторая запись на тот же тест не проходит
    даже в обход upsert_haplogroup(), напрямую через SQL."""
    gns = _make_test(conn)
    ts = repo.now_iso()
    conn.execute(
        "INSERT INTO haplogroups (test_number, created_at, updated_at) VALUES (?, ?, ?)",
        (gns, ts, ts),
    )
    with pytest.raises(sqlite3.IntegrityError):
        conn.execute(
            "INSERT INTO haplogroups (test_number, created_at, updated_at) VALUES (?, ?, ?)",
            (gns, ts, ts),
        )


def test_db_level_fk_blocks_deleting_referenced_test(conn):
    """Тест с записью гаплогруппы нельзя удалить из tests напрямую — FK +
    PRAGMA foreign_keys=ON останавливают это на уровне БД (приложение и так
    никогда не удаляет тесты физически, только архивирует)."""
    gns = _make_test(conn)
    hrepo.upsert_haplogroup(conn, gns, {"y_dna": "R1a"})
    test_id = conn.execute("SELECT test_id FROM tests WHERE gns_number=?", (gns,)).fetchone()["test_id"]
    with pytest.raises(sqlite3.IntegrityError):
        conn.execute("DELETE FROM tests WHERE test_id = ?", (test_id,))


def test_snp_search_matches_substring_mid_string(conn):
    """Поиск по SNP — contains, а не точное совпадение (напр. 'PH908' находится
    и внутри 'I-PH908', и внутри 'I2a1...CTS10228 > Y3120 > PH908')."""
    gns1 = _make_test(conn, customer_name="Носитель PH908")
    gns2 = _make_test(conn, customer_name="Носитель YP417")
    hrepo.upsert_haplogroup(conn, gns1, {"snp_issued": "I-PH908"})
    hrepo.upsert_haplogroup(conn, gns2, {"snp_issued": "R-YP417"})

    assert {r["test_number"] for r in hrepo.list_haplogroups(conn, snp_issued="PH908")} == {gns1}
    assert {r["test_number"] for r in hrepo.list_haplogroups(conn, snp_issued="YP417")} == {gns2}
    # частичный фрагмент без начала/конца строки тоже должен находиться
    assert {r["test_number"] for r in hrepo.list_haplogroups(conn, snp_issued="H908")} == {gns1}


def test_project_rename_keeps_haplogroup_linkage_via_id_not_name(conn):
    """Связь идёт по project_id, а не по названию — переименование проекта
    не рвёт привязку существующих записей."""
    bashkir = conn.execute("SELECT id FROM projects WHERE name = 'Башкирский проект'").fetchone()["id"]
    gns = _make_test(conn)
    hrepo.upsert_haplogroup(conn, gns, {"project_id": bashkir})

    prepo.rename_project(conn, bashkir, "Башкирский ДНК-проект")

    row = hrepo.list_haplogroups(conn, test_number=gns)[0]
    assert row["project_id"] == bashkir
    assert prepo.get_project_path(conn, row["project_id"]) == "Этнопроекты › Башкирский ДНК-проект"


def test_archiving_project_does_not_hide_or_orphan_existing_records(conn):
    """Архивирование проекта не удаляет и не скрывает уже привязанные к нему
    записи гаплогрупп — только помечает сам проект is_archived=1 и убирает
    его из списка для выбора в новых записях."""
    tatar = conn.execute("SELECT id FROM projects WHERE name = 'Татарский проект'").fetchone()["id"]
    gns = _make_test(conn)
    hrepo.upsert_haplogroup(conn, gns, {"project_id": tatar})

    prepo.archive_project(conn, tatar)

    # запись гаплогруппы по-прежнему на месте и привязана к проекту
    row = hrepo.list_haplogroups(conn, test_number=gns)[0]
    assert row["project_id"] == tatar
    row2 = conn.execute("SELECT is_archived FROM projects WHERE id=?", (tatar,)).fetchone()
    assert row2["is_archived"] == 1

    # но архивный проект больше не предлагается для выбора в новых записях
    active_ids = {pid for pid, _n, _d, _a in prepo.flatten_tree(prepo.build_tree(conn, include_archived=False))}
    assert tatar not in active_ids
    # хотя в полном дереве (для справочника/фильтра) он всё ещё виден
    all_ids = {pid for pid, _n, _d, _a in prepo.flatten_tree(prepo.build_tree(conn, include_archived=True))}
    assert tatar in all_ids


def test_color_logic_deeper_nevgen_snp_still_found_in_semargl_chain_is_green(conn):
    """Ревью-кейс: NevGen называет более глубокий/более неглубокий SNP той же
    цепочки, что и Semargl — раз терминальный SNP NevGen встречается где-то
    в строке Semargl, это согласованная предикция (зелёный), а не жёлтый."""
    semargl = "I2a1b3a1a1c CTS10228 > Y3120 > PH908"
    assert hlogic.calculate_haplo_color("I-PH908", semargl, "I2a") == hlogic.HaploColor.GREEN
    assert hlogic.calculate_haplo_color("I-CTS10228", semargl, "I2a") == hlogic.HaploColor.GREEN


def test_final_haplogroup_reserved_field_present_and_hidden(conn):
    cols = {r[1] for r in conn.execute("PRAGMA table_info(haplogroups)").fetchall()}
    assert "final_haplogroup" in cols
    assert "final_haplogroup" in hrepo.RESERVED_FIELDS
    assert "final_haplogroup" not in hrepo.EDITABLE_FIELDS

    gns = _make_test(conn)
    hrepo.upsert_haplogroup(conn, gns, {"final_haplogroup": "I-PH908"})
    row = conn.execute("SELECT final_haplogroup FROM haplogroups WHERE test_number=?", (gns,)).fetchone()
    assert row["final_haplogroup"] == "I-PH908"
