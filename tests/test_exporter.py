"""Тесты экспорта заказов/тестов в Excel (задача 4 финального этапа)."""
import os
import sys
from datetime import date

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

import openpyxl
import pytest

from app import db as dbmod
from app import repository as repo
from app import exporter


@pytest.fixture()
def conn(tmp_path):
    c = dbmod.connect(tmp_path / "test.db")
    dbmod.init_db(c)
    yield c
    c.close()


def test_export_creates_valid_xlsx_with_header_and_rows(conn, tmp_path):
    order = repo.NewOrderInput(order_no="OZ-500", customer_name="Экспорт Тестович",
                                test_type_codes=["Y50", "MTDNA"], report_added=True,
                                sample_received_at=date(2026, 1, 1))
    repo.create_order_with_tests(conn, order)

    out_path = tmp_path / "export.xlsx"
    count = exporter.export_orders_to_excel(conn, str(out_path))
    assert count == 2
    assert out_path.exists()

    wb = openpyxl.load_workbook(out_path)
    ws = wb.active
    header = [c.value for c in ws[1]]
    assert "GNS" in header
    assert "ФИО" in header
    assert ws.max_row == 3  # заголовок + 2 теста


def test_export_respects_search_filter(conn, tmp_path):
    repo.create_order_with_tests(conn, repo.NewOrderInput(
        order_no="OZ-1", customer_name="Иванов", test_type_codes=["Y50"],
        sample_received_at=date(2026, 1, 1)))
    repo.create_order_with_tests(conn, repo.NewOrderInput(
        order_no="OZ-2", customer_name="Петров", test_type_codes=["Y50"],
        sample_received_at=date(2026, 1, 1)))

    out_path = tmp_path / "filtered.xlsx"
    count = exporter.export_orders_to_excel(conn, str(out_path), search="Иванов")
    assert count == 1
    wb = openpyxl.load_workbook(out_path)
    ws = wb.active
    fio_col = [c.value for c in ws[1]].index("ФИО")
    assert ws.cell(row=2, column=fio_col + 1).value == "Иванов"


def test_export_empty_db_produces_header_only(conn, tmp_path):
    out_path = tmp_path / "empty.xlsx"
    count = exporter.export_orders_to_excel(conn, str(out_path))
    assert count == 0
    wb = openpyxl.load_workbook(out_path)
    assert wb.active.max_row == 1  # только заголовок


def test_export_includes_comments_and_results(conn, tmp_path):
    order = repo.NewOrderInput(customer_name="С Результатом", test_type_codes=["Y50"],
                                sample_received_at=date(2026, 1, 1))
    order_id, test_ids = repo.create_order_with_tests(conn, order)
    repo.update_test_fields(conn, test_ids[0], {"result_y": "R1a-M198", "comments": "особый случай"})

    out_path = tmp_path / "results.xlsx"
    exporter.export_orders_to_excel(conn, str(out_path))
    wb = openpyxl.load_workbook(out_path)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    y_col = headers.index("Результат Y") + 1
    comment_col = headers.index("Комментарий") + 1
    assert ws.cell(row=2, column=y_col).value == "R1a-M198"
    assert ws.cell(row=2, column=comment_col).value == "особый случай"
